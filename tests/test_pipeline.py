"""בדיקות למשיכת ההגשות לתיקיית OneDrive. מריצים: python -m pytest"""
from __future__ import annotations

import http.client
import io
import json
import threading
import zipfile
from datetime import datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, quote

import pytest
from openpyxl import load_workbook

from pipeline import config, service, settings as cfg
from pipeline.demo import homework_pdf, moodle_zip
from pipeline.gradebook import GradeBook
from pipeline.moodle_api import Moodle
from pipeline.moodle_zip import read_zip
from pipeline.pull import pull_zip
from pipeline.workspace import Workspace

COURSE, HW = "אלגברה לינארית 1", "תרגיל בית 1"


@pytest.fixture
def root(tmp_path, monkeypatch):
    r = tmp_path / "OneDrive - BGU" / config.FOLDER_NAME
    r.mkdir(parents=True)
    monkeypatch.setenv("ASSIGNMENT_CHECKER_ROOT", str(r))
    monkeypatch.delenv("MOODLE_TOKEN", raising=False)
    return r


def rows(path: Path) -> list[dict]:
    ws = load_workbook(path)["ציונים"]
    head = [c.value for c in ws[1]]
    return [dict(zip(head, (c.value for c in r))) for r in ws.iter_rows(min_row=2)]


class _NoUtf8Flag(zipfile.ZipInfo):
    """ZIP ישן: השם נשמר בבתים של UTF-8, אבל בלי הסימון שאומר שזה UTF-8."""
    def _encodeFilenameFlags(self):
        return self.filename.encode("utf-8"), self.flag_bits & ~0x800


def make_zip(entries: dict[str, bytes], utf8: bool = True) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for name, data in entries.items():
            cls = zipfile.ZipInfo if utf8 else _NoUtf8Flag
            zf.writestr(cls(name, date_time=(2026, 10, 20, 12, 0, 0)), data)
    return buf.getvalue()


# ---------- ZIP של מודל ----------
def test_read_zip_folders_flat_and_wrapper(tmp_path):
    pdf = homework_pdf("נועה כהן", HW)
    z = tmp_path / "x.zip"
    z.write_bytes(make_zip({
        "נועה כהן_111_assignsubmission_file/hw.pdf": pdf,
        "איתי לוי_222_assignsubmission_file_hw.pdf": pdf,
        "COURSE-HW1-9/מאיה מזרחי_333_assignsubmission_file/part 1.pdf": pdf,
        "מאיה מזרחי_333_assignsubmission_onlinetext/onlinetext.html": b"<p>hi</p>",
        "__MACOSX/junk": b"x",
    }))
    subs = {s.moodle_id: s for s in read_zip(z)}
    assert set(subs) == {"111", "222", "333"}
    assert subs["111"].name == "נועה כהן" and subs["111"].files[0].name == "hw.pdf"
    assert subs["222"].files[0].name == "hw.pdf"
    assert sorted(f.name for f in subs["333"].files) == ["onlinetext.html", "part 1.pdf"]
    assert subs["111"].submitted_at == datetime(2026, 10, 20, 12, 0)


def test_read_zip_hebrew_names_without_utf8_flag(tmp_path):
    z = tmp_path / "old.zip"
    z.write_bytes(make_zip({"שירה ביטון_444_assignsubmission_file_hw.pdf": b"%PDF"}, utf8=False))
    (s,) = read_zip(z)
    assert s.name == "שירה ביטון"


# ---------- קליטה לתיקייה ----------
def test_pull_creates_folders_excel_and_flags_problems(root):
    ws = Workspace(root, COURSE, HW)
    ws.ensure()
    z = ws.inbox / "ALG1-תרגיל בית 1-77.zip"
    z.write_bytes(moodle_zip(HW, count=5, seed=1))
    rep = service.pull(root, COURSE, HW)

    assert rep.result.total == 5 and rep.result.new == 5
    assert len(rep.result.problems) == 1  # הדגמה מכניסה קובץ ריק אחד
    folders = sorted(p.name for p in ws.submissions.iterdir())
    assert len(folders) == 5 and all(f.endswith(")") for f in folders)
    assert not list(ws.inbox.glob("*.zip"))  # ה-ZIP עבר ל"נקלטו"
    assert len(list((ws.inbox / config.INBOX_DONE).glob("*.zip"))) == 1

    table = rows(ws.grades_path)
    assert len(table) == 5
    statuses = sorted(r["מצב"] for r in table)
    assert statuses.count("ממתינה לבדיקה") == 4 and statuses.count("דורשת טיפול") == 1
    assert all(r["עמודים"] >= 1 for r in table if r["מצב"] == "ממתינה לבדיקה")
    assert load_workbook(ws.grades_path)["ציונים"].sheet_view.rightToLeft
    assert (ws.dir / config.LOG_FILE).read_text(encoding="utf-8").count("נקלט") == 1
    assert (root / service.README_NAME).exists() and (root / cfg.SETTINGS_FILE).exists()


def test_second_pull_keeps_team_edits_and_detects_resubmission(root):
    ws = Workspace(root, COURSE, HW)
    ws.ensure()
    data = make_zip({"נועה כהן_111_assignsubmission_file/hw.pdf": homework_pdf("נועה כהן", HW),
                     "איתי לוי_222_assignsubmission_file/hw.pdf": homework_pdf("איתי לוי", HW)})
    (ws.inbox / "a.zip").write_bytes(data)
    service.pull(root, COURSE, HW)

    # הצוות בודק ומאשר בקובץ
    book = GradeBook(ws.grades_path)
    for mid, r in book.rows().items():
        book.ws.cell(row=r, column=book.cols.index("מצב") + 1, value="נבדקה")
        book.ws.cell(row=r, column=book.cols.index("ציון") + 1, value=90)
        book.ws.cell(row=r, column=book.cols.index("משוב") + 1, value="יפה מאוד")
        book.ws.cell(row=r, column=book.cols.index("מאושר") + 1, value="כן")
    book.save()

    # אותו ZIP שוב: כלום לא משתנה
    (ws.inbox / "a.zip").write_bytes(data)
    rep = service.pull(root, COURSE, HW)
    assert rep.result.same == 2 and rep.result.new == 0
    assert all(r["ציון"] == 90 and r["משוב"] == "יפה מאוד" and r["מצב"] == "נבדקה" for r in rows(ws.grades_path))

    # נועה הגישה מחדש
    (ws.inbox / "b.zip").write_bytes(make_zip({"נועה כהן_111_assignsubmission_file/hw.pdf": homework_pdf("נועה כהן", HW, pages=3)}))
    rep = service.pull(root, COURSE, HW)
    assert rep.result.updated == 1
    noa = next(r for r in rows(ws.grades_path) if r["מזהה במודל"] == "111")
    assert noa["מצב"] == "ממתינה לבדיקה" and "הוגשה מחדש" in noa["הערות מערכת"]
    assert noa["ציון"] == 90  # הציון הקודם נשאר, הצוות מחליט מה לעשות
    itai = next(r for r in rows(ws.grades_path) if r["מזהה במודל"] == "222")
    assert itai["מצב"] == "נבדקה"


def test_late_submission_is_marked(root):
    s = service.init(root)
    s.course(COURSE).assignment(HW).due = "2026-10-20T10:00"
    cfg.save(root, s)
    ws = Workspace(root, COURSE, HW)
    (ws.inbox / "a.zip").write_bytes(make_zip({"נועה כהן_111_assignsubmission_file/hw.pdf": homework_pdf("נועה כהן", HW)}))
    service.pull(root, COURSE, HW)
    (r,) = rows(ws.grades_path)
    assert "באיחור" in r["הערות מערכת"] and r["מצב"] == "ממתינה לבדיקה"


# ---------- קליטה ראשית: ניתוב לפי שם הקובץ ----------
def test_root_inbox_routes_by_moodle_zip_name(root):
    service.init(root)
    (root / config.INBOX).mkdir(exist_ok=True)
    (root / config.INBOX / "ALG1-תרגיל בית 3-555.zip").write_bytes(moodle_zip("תרגיל בית 3", count=3, seed=2, empty_one=False))
    (root / config.INBOX / "ALG1-תרגיל חדש-556.zip").write_bytes(moodle_zip("תרגיל חדש", count=2, seed=3, empty_one=False))
    rep = service.pull(root, COURSE, "תרגיל בית 3")

    assert rep.result.new == 3
    assert any("תרגיל חדש" in n for n in rep.notes)
    assert len(list(Workspace(root, COURSE, "תרגיל חדש").submissions.iterdir())) == 2
    s = cfg.load(root)
    assert s.course(COURSE).assignment("תרגיל חדש").moodle_cmid == 556
    assert s.course(COURSE).assignment("תרגיל חדש").last_pulled  # גם היא מסומנת כנמשכה
    assert rep.routed == [{"course": COURSE, "assignment": "תרגיל חדש", "total": 2}]
    assert not list((root / config.INBOX).glob("*.zip"))


def test_no_sources_explains_what_to_do(root):
    rep = service.pull(root, COURSE, HW)
    assert rep.result.total == 0
    assert any("ZIP" in n for n in rep.notes)


def test_find_root_explains_when_missing(tmp_path, monkeypatch):
    monkeypatch.delenv("ASSIGNMENT_CHECKER_ROOT", raising=False)
    monkeypatch.setattr(config, "onedrive_roots", lambda: [tmp_path])
    with pytest.raises(config.FolderNotFound) as e:
        config.find_root()
    assert "Add shortcut to My files" in str(e.value)
    (tmp_path / "Shared" / config.FOLDER_NAME).mkdir(parents=True)
    assert config.find_root() == tmp_path / "Shared" / config.FOLDER_NAME


# ---------- מודל עם טוקן (שרת מודל מדומה) ----------
class FakeMoodle(BaseHTTPRequestHandler):
    pdf = b""

    def log_message(self, *a):
        pass

    def do_POST(self):
        form = {k: v[0] for k, v in parse_qs(self.rfile.read(int(self.headers["Content-Length"])).decode()).items()}
        assert form["wstoken"] == "tok"
        fn = form["wsfunction"]
        if fn == "mod_assign_get_assignments":
            out = {"courses": [{"id": 9, "assignments": [{"id": 301, "cmid": 9001, "name": "תרגיל בית 1", "duedate": 1792540740},
                                                          {"id": 302, "cmid": 9002, "name": "תרגיל בית 2", "duedate": 0}]}]}
        elif fn == "mod_assign_get_submissions":
            base = f"http://127.0.0.1:{self.server.server_port}/webservice/pluginfile.php"
            out = {"assignments": [{"assignmentid": 301, "submissions": [
                {"userid": 5, "status": "submitted", "timemodified": 1792500000, "plugins": [
                    {"type": "file", "fileareas": [{"area": "submission_files", "files": [{"filename": "hw.pdf", "fileurl": f"{base}/1/hw.pdf"}]}]}]},
                {"userid": 6, "status": "new", "timemodified": 0, "plugins": []},
            ]}]}
        elif fn == "core_user_get_users_by_field":
            out = [{"id": 5, "fullname": "נועה כהן", "email": "noa@post.bgu.ac.il", "idnumber": "123456789"}]
        else:
            out = {"exception": "x", "message": "unknown"}
        body = json.dumps(out).encode()
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self):
        assert "token=tok" in self.path
        self.send_response(200)
        self.end_headers()
        self.wfile.write(FakeMoodle.pdf)


def test_pull_from_moodle_with_token(root, monkeypatch):
    FakeMoodle.pdf = homework_pdf("נועה כהן", HW)
    srv = ThreadingHTTPServer(("127.0.0.1", 0), FakeMoodle)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        s = service.init(root)
        s.moodle_url = f"http://127.0.0.1:{srv.server_port}"
        s.course(COURSE).moodle_course_id = 9
        cfg.save(root, s)
        monkeypatch.setenv("MOODLE_TOKEN", "tok")

        rep = service.pull(root, COURSE, HW)
        assert "מודל" in rep.source and rep.result.new == 1
        (r,) = rows(Workspace(root, COURSE, HW).grades_path)
        assert r["שם מלא"] == "נועה כהן" and r["ת״ז"] == "123456789" and r["מייל"] == "noa@post.bgu.ac.il"
        s = cfg.load(root)
        assert s.course(COURSE).assignment(HW).moodle_assign_id == 301
        assert s.course(COURSE).assignment(HW).due  # תאריך ההגשה הגיע ממודל
        assert Workspace(root, COURSE, "תרגיל בית 2").dir.is_dir()  # מטלה ממודל קיבלה תיקייה לבד
    finally:
        srv.shutdown()


def test_moodle_error_is_readable():
    m = Moodle("http://127.0.0.1:9", "tok", timeout=2)
    with pytest.raises(Exception) as e:
        m.assignments(1)
    assert "אין חיבור למודל" in str(e.value)


# ---------- השרת המקומי ----------
@pytest.fixture
def server(root):
    from pipeline.server import Handler
    srv = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    yield srv.server_port
    srv.shutdown()


def call(port, method, path, body=None, headers=None):
    conn = http.client.HTTPConnection("127.0.0.1", port, timeout=30)
    h = {"Host": f"localhost:{port}", **(headers or {})}
    data = json.dumps(body).encode() if body is not None else None
    if data:
        h["Content-Type"] = "application/json"
    conn.request(method, path, body=data, headers=h)
    r = conn.getresponse()
    raw = r.read()
    return r.status, (json.loads(raw) if r.getheader("Content-Type", "").startswith("application/json") else raw)


def test_server_health_state_pull_and_files(root, server):
    st, h = call(server, "GET", "/api/health")
    assert st == 200 and h["root"] == str(root)

    ws = Workspace(root, COURSE, HW)
    ws.ensure()
    (ws.inbox / "a.zip").write_bytes(moodle_zip(HW, count=3, seed=4))
    st, rep = call(server, "POST", "/api/pull", {"course": COURSE, "assignment": HW}, {"X-Assignment-Checker": "1"})
    assert st == 200 and rep["result"]["new"] == 3

    st, state = call(server, "GET", "/api/state")
    a = next(x for x in state["assignments"] if x["title"] == HW)
    assert len(a["subs"]) == 3 and a["pulledAt"]
    assert {s["status"] for s in a["subs"]} == {"pulled", "error"}
    f = next(s for s in a["subs"] if s["status"] == "pulled")["files"][0]
    st, data = call(server, "GET", f"/api/file?path={quote(f['path'])}")
    assert st == 200 and data[:4] == b"%PDF"

    st, index = call(server, "GET", "/")
    assert st == 200 and b"<html" in index


def test_server_refuses_foreign_requests(root, server):
    # בלי הכותרת המיוחדת: אתר אחר לא יכול להפעיל משיכה
    st, _ = call(server, "POST", "/api/pull", {"course": COURSE, "assignment": HW})
    assert st == 403
    # מקור זר
    st, _ = call(server, "POST", "/api/pull", {"course": COURSE, "assignment": HW}, {"X-Assignment-Checker": "1", "Origin": "https://evil.example"})
    assert st == 403
    # Host זר (DNS rebinding)
    st, _ = call(server, "GET", "/api/state", headers={"Host": "evil.example"})
    assert st == 403
    # קבצים מחוץ לתיקיית הקורס
    st, _ = call(server, "GET", "/api/file?path=../../secret.txt")
    assert st == 404
